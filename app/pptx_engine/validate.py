"""Post-build checks run before a deck is offered for download."""

from __future__ import annotations

import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation as PptxPresentation
from pptx.util import Emu

from .textfit import text_height_in


@dataclass
class ValidationReport:
    ok: bool
    slide_count: int
    pictures: int
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


def validate_pptx(path: Path, expected_slides: int | None = None) -> ValidationReport:
    errors: list[str] = []
    warnings: list[str] = []

    if not path.is_file() or path.stat().st_size == 0:
        return ValidationReport(False, 0, 0, ["file missing or empty"])
    try:
        with zipfile.ZipFile(path) as zf:
            bad = zf.testzip()
            if bad:
                errors.append(f"corrupted zip member: {bad}")
    except zipfile.BadZipFile:
        return ValidationReport(False, 0, 0, ["not a valid pptx (zip) file"])

    try:
        prs = PptxPresentation(str(path))
    except Exception as exc:  # python-pptx raises many types on broken XML
        return ValidationReport(False, 0, 0, [f"cannot open with python-pptx: {exc}"])

    if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
        errors.append("slide size is not 16:9")

    sw, sh = prs.slide_width, prs.slide_height
    pictures = 0
    for idx, slide in enumerate(prs.slides, start=1):
        texts = [sh_.text_frame.text for sh_ in slide.shapes if sh_.has_text_frame and sh_.text_frame.text.strip()]
        if not texts:
            errors.append(f"slide {idx} has no text")
        for shp in slide.shapes:
            if shp.shape_type == 13:  # PICTURE
                pictures += 1
            if shp.has_text_frame and shp.text_frame.text.strip() and shp.width and shp.height:
                # Re-check overflow on the saved file using the actual font sizes.
                paras = [p.text for p in shp.text_frame.paragraphs if p.text.strip()]
                sizes = [r.font.size.pt for p in shp.text_frame.paragraphs for r in p.runs if r.font.size]
                if sizes:
                    bold = any(r.font.bold for p in shp.text_frame.paragraphs for r in p.runs)
                    need = text_height_in(paras, Emu(shp.width).inches, max(sizes), bold)
                    if need > Emu(shp.height).inches * 1.08:
                        warnings.append(f"slide {idx}: text may overflow ('{paras[0][:30]}…')")
            # Decorative shapes may bleed off-canvas on purpose; flag only text boxes.
            if shp.has_text_frame and shp.text_frame.text.strip():
                if shp.left < 0 or shp.top < 0 or shp.left + shp.width > sw * 1.001 or shp.top + shp.height > sh * 1.001:
                    errors.append(f"slide {idx}: text box outside slide")

    count = len(prs.slides)
    if count == 0:
        errors.append("presentation has no slides")
    if expected_slides is not None and count != expected_slides:
        errors.append(f"expected {expected_slides} slides, found {count}")
    return ValidationReport(not errors, count, pictures, errors, warnings)
