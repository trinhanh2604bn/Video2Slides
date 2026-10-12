"""In-memory stand-in for Google OAuth, Drive and Slides REST endpoints (httpx.MockTransport).

Conversion is simulated by reading the uploaded .pptx with python-pptx, so the slide
count, texts, pictures and notes the fake "Google Slides" exposes come from the real deck.
"""

from __future__ import annotations

import io
import json
import re
from urllib.parse import parse_qs

import httpx
from pptx import Presentation as PptxPresentation

PNG = b"\x89PNG\r\n\x1a\n" + b"0" * 64


class FakeGoogle:
    def __init__(self):
        self.files: dict[str, dict] = {}
        self.decks: dict[str, dict] = {}
        self.uploads: dict[str, dict] = {}
        self.calls: list[tuple[str, str]] = []
        self.scope = "https://www.googleapis.com/auth/drive.file"
        self.refresh_error: str | None = None
        self.drop_notes: set[int] = set()        # slide indexes that lose their notes in conversion
        self.fail_queue: dict[str, list[int]] = {}  # url substring -> statuses returned first
        self.disabled_api = False
        self.token_posts: list[dict] = []
        self.revoked: list[str] = []
        self._n = 0

    def client(self) -> httpx.Client:
        return httpx.Client(transport=httpx.MockTransport(self.handler))

    def count(self, method: str, needle: str) -> int:
        return sum(1 for m, u in self.calls if m == method and needle in u)

    # ------------------------------------------------------------------ routing
    def handler(self, request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        self.calls.append((request.method, url))
        for needle, statuses in self.fail_queue.items():
            if needle in url and statuses:
                return httpx.Response(statuses.pop(0), json={"error": {"message": "busy"}})
        host = request.url.host
        if host == "oauth2.googleapis.com":
            return self._oauth(request)
        if host.endswith("googleusercontent.com"):
            return httpx.Response(200, content=PNG)
        if not request.headers.get("authorization", "").startswith("Bearer at-"):
            return httpx.Response(401, json={"error": {"message": "unauthenticated"}})
        if self.disabled_api:
            return httpx.Response(403, json={"error": {"message": "Google Slides API has not been used in project 1 "
                                                                  "before or it is disabled.",
                                                       "details": [{"reason": "SERVICE_DISABLED"}]}})
        path = request.url.path
        if path == "/drive/v3/about":
            return httpx.Response(200, json={"user": {"emailAddress": "test@example.com", "displayName": "Test"}})
        if path == "/upload/drive/v3/files" and request.method == "POST":
            return self._start_upload(request)
        if path == "/upload/drive/v3/files" and request.method == "PUT":
            return self._finish_upload(request)
        if path == "/drive/v3/files" and request.method == "GET":
            return self._list(request)
        if m := re.fullmatch(r"/drive/v3/files/([\w-]+)", path):
            f = self.files.get(m.group(1))
            if not f:
                return httpx.Response(404, json={"error": {"message": "File not found"}})
            return httpx.Response(200, json={"id": m.group(1), "trashed": f.get("trashed", False)})
        if m := re.fullmatch(r"/v1/presentations/([\w-]+):batchUpdate", path):
            return self._batch(m.group(1), json.loads(request.content))
        if m := re.fullmatch(r"/v1/presentations/([\w-]+)/pages/([\w-]+)/thumbnail", path):
            return httpx.Response(200, json={"contentUrl": "https://lh3.googleusercontent.com/thumb-x"})
        if m := re.fullmatch(r"/v1/presentations/([\w-]+)", path):
            deck = self.decks.get(m.group(1))
            if not deck:
                return httpx.Response(404, json={"error": {"message": "not found"}})
            return httpx.Response(200, json=self.doc(m.group(1)))
        return httpx.Response(404, json={"error": {"message": f"unknown {path}"}})

    def _oauth(self, request: httpx.Request) -> httpx.Response:
        form = {k: v[0] for k, v in parse_qs(request.content.decode()).items()}
        if request.url.path == "/revoke":
            self.revoked.append(form.get("token", ""))
            return httpx.Response(200)
        self.token_posts.append(form)
        if form.get("grant_type") == "authorization_code":
            return httpx.Response(200, json={"access_token": "at-1", "expires_in": 3599, "token_type": "Bearer",
                                             "refresh_token": "rt-SECRET-123", "scope": self.scope})
        if form.get("grant_type") == "refresh_token":
            if self.refresh_error:
                return httpx.Response(400, json={"error": self.refresh_error})
            return httpx.Response(200, json={"access_token": "at-2", "expires_in": 3599, "scope": self.scope})
        return httpx.Response(400, json={"error": "unsupported_grant_type"})

    # ------------------------------------------------------------------ drive
    def _start_upload(self, request: httpx.Request) -> httpx.Response:
        self._n += 1
        uid = f"up{self._n}"
        self.uploads[uid] = json.loads(request.content)
        return httpx.Response(200, headers={"Location": f"https://www.googleapis.com/upload/drive/v3/files"
                                                        f"?uploadType=resumable&upload_id={uid}"})

    def _finish_upload(self, request: httpx.Request) -> httpx.Response:
        uid = request.url.params["upload_id"]
        meta = self.uploads.pop(uid)
        self._n += 1
        fid = f"pres{self._n}"
        self.files[fid] = {**meta, "trashed": False}
        prs = PptxPresentation(io.BytesIO(request.content))
        slides = []
        for i, sl in enumerate(prs.slides):
            texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            pics = sum(1 for sh in sl.shapes if sh.shape_type == 13)
            notes = sl.notes_slide.notes_text_frame.text if sl.has_notes_slide else ""
            slides.append({"objectId": f"{fid}_s{i}", "texts": texts, "images": pics,
                           "notes": None if i in self.drop_notes else notes})
        self.decks[fid] = {"title": meta["name"], "slides": slides}
        return httpx.Response(200, json={"id": fid})

    def _list(self, request: httpx.Request) -> httpx.Response:
        q = request.url.params.get("q", "")
        want = dict(re.findall(r"key='(\w+)' and value='([^']*)'", q))
        out = [{"id": fid, "name": f["name"]} for fid, f in self.files.items()
               if not f.get("trashed") and all(f["appProperties"].get(k) == v for k, v in want.items())]
        return httpx.Response(200, json={"files": out})

    # ------------------------------------------------------------------ slides
    def doc(self, fid: str) -> dict:
        deck = self.decks[fid]
        slides = []
        for s in deck["slides"]:
            els = [{"objectId": f"{s['objectId']}_t{k}", "shape": {"text": {"textElements": [
                {"textRun": {"content": t + "\n"}}]}}} for k, t in enumerate(s["texts"])]
            els += [{"objectId": f"{s['objectId']}_i{k}", "image": {"contentUrl": "https://x"}} for k in range(s["images"])]
            nid = f"{s['objectId']}_notes"
            notes_els = [] if s["notes"] is None else [
                {"objectId": nid, "shape": {"text": {"textElements": [{"textRun": {"content": s["notes"] + "\n"}}]}}}]
            slides.append({"objectId": s["objectId"], "pageElements": els, "slideProperties": {"notesPage": {
                "notesProperties": {"speakerNotesObjectId": nid}, "pageElements": notes_els}}})
        return {"presentationId": fid, "title": deck["title"], "slides": slides,
                "pageSize": {"width": {"magnitude": 12192000, "unit": "EMU"},
                             "height": {"magnitude": 6858000, "unit": "EMU"}}}

    def _batch(self, fid: str, body: dict) -> httpx.Response:
        by_notes = {f"{s['objectId']}_notes": s for s in self.decks[fid]["slides"]}
        for req in body["requests"]:
            if "deleteText" in req:
                s = by_notes[req["deleteText"]["objectId"]]
                if not s["notes"]:
                    return httpx.Response(400, json={"error": {"message": "deleteText on empty shape"}})
                s["notes"] = ""
            elif "insertText" in req:
                s = by_notes[req["insertText"]["objectId"]]
                s["notes"] = (s["notes"] or "") + req["insertText"]["text"]
        return httpx.Response(200, json={"replies": [{} for _ in body["requests"]]})
