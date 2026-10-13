"""Cloud pipeline: upload -> FFmpeg/Whisper on the Modal media worker -> routed LLM -> PPTX.

MOCKED: tests/fake_modal.py stands in for both Modal workers and Gemini is an httpx MockTransport. Nothing here
proves GPU execution; the live check is scripts/modal_media_live_check.py.
"""

import json
import time
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation as PptxFile

from app import jobs, main, pipeline
from app.ai import gemini
from app.ai.router import CircuitBreaker, RoutingProvider
from tests.fake_modal import FakeModalLLM, FakeModalMedia, use_settings
from tests.conftest import needs_ffmpeg
from tests.test_failures import _wait


@pytest.fixture
def cloud(monkeypatch, tmp_path):
    media = FakeModalMedia()
    llm = FakeModalLLM()
    use_settings(monkeypatch, media_processing="modal", modal_media_worker_url="https://media.modal.test",
                 modal_poll_interval_sec=0, gemini_api_key="", modal_media_upload_chunk_mb=1)
    monkeypatch.setattr(pipeline, "ModalMediaClient", media.client_factory())
    # Local FFmpeg / Whisper must not be used in the cloud path.
    boom = lambda *a, **k: (_ for _ in ()).throw(AssertionError("local media processing used"))  # noqa: E731
    for name in ("transcribe", "extract_audio", "extract_frames"):
        monkeypatch.setattr(pipeline, name, boom)
    monkeypatch.setattr(main, "ffmpeg_available", lambda: {"ffmpeg": None, "ffprobe": None})
    monkeypatch.setattr(pipeline, "get_provider", lambda: RoutingProvider(
        circuit=CircuitBreaker(tmp_path / "circuit.json"), modal_factory=llm.provider))
    with TestClient(main.app) as c:
        yield c, media, llm


_VIDEO: dict = {}


def _mp4() -> bytes:
    """A real (synthetic) MP4: PyAV validates uploads and copies the audio track when FFmpeg is absent."""
    if "data" not in _VIDEO:
        import tempfile

        from tests.conftest import make_video
        path = Path(tempfile.mkdtemp()) / "v.mp4"
        _VIDEO["data"] = make_video(path, seconds=6).read_bytes()
    return _VIDEO["data"]


def _submit(client, data: bytes | None = None, **form) -> str:
    resp = client.post("/api/jobs", data={"slide_count": "6", **form},
                       files={"file": ("bai-giang.mp4", data or _mp4(), "video/mp4")})
    assert resp.status_code == 202, resp.text
    return resp.json()["job_id"]


@needs_ffmpeg
def test_full_cloud_pipeline_without_local_ffmpeg_whisper_or_gemini_key(cloud):
    client, media, llm = cloud
    job_id = _submit(client)
    job = _wait(client, job_id)
    assert job["status"] == "done", job["error"]
    d = jobs.job_dir(job_id)
    # chunked upload reached the worker intact, temp files deleted on Modal after transcription
    upload = next(iter(media.uploads.values()))
    sent = sum(len(c) for c in upload["chunks"].values())
    assert media.deleted and sent < len(_mp4()) / 3  # only the remuxed audio track was uploaded
    assert media.preprocess_params[-1]["audio_only"] is True
    assert not (d / "remote_media.json").exists() and not (d / "audio_upload.m4a").exists()
    meta = json.loads((d / "transcript_meta.json").read_text(encoding="utf-8"))
    assert meta["location"] == "modal" and meta["compute_type"] == "float16" and meta["model"] == "small"
    assert media.transcribe_params[-1]["batch_size"] == 16 and media.transcribe_params[-1]["word_timestamps"]
    assert meta["processing"]["whisper"]["device"] == "cuda"
    assert job["media"]["duration"] == 6.0 and len(json.loads((d / "frames.json").read_text())) >= 1
    result = job["result"]
    assert result["provider"] == "modal_ollama" and result["llm_route"]["fallback_reason"] == "missing_key"
    assert result["media_processing"]["ffmpeg"]["location"] == "modal-cpu"
    out = json.loads((d / "outline_meta.json").read_text(encoding="utf-8"))
    assert out["provider"] == "modal_ollama" and out["transcript_key"] == meta["key"]

    # Notes belong to the right slides (slide i's notes contain slide i's speaker notes)
    pres, _ = pipeline.load_outline(d)
    deck = PptxFile(str(pipeline.output_path(d)))
    assert len(deck.slides) == len(pres.slides)
    for slide, src in zip(deck.slides, pres.slides):
        notes = slide.notes_slide.notes_text_frame.text
        assert src.speaker_notes.strip()[:20] in notes


@needs_ffmpeg
def test_video_upload_mode_still_works(cloud, monkeypatch):
    client, media, _ = cloud
    use_settings(monkeypatch, media_processing="modal", modal_media_worker_url="https://media.modal.test",
                 modal_poll_interval_sec=0, gemini_api_key="", modal_upload_audio_only=False)
    job = _wait(client, _submit(client))
    assert job["status"] == "done", job["error"]
    assert media.preprocess_params[-1]["audio_only"] is False and job["media"]["width"] == 640


@needs_ffmpeg
def test_gemini_failure_falls_back_without_rerunning_transcription(cloud, monkeypatch, tmp_path):
    client, media, llm = cloud
    monkeypatch.setattr(gemini.time, "sleep", lambda s: None)
    use_settings(monkeypatch, media_processing="modal", modal_media_worker_url="https://media.modal.test",
                 modal_poll_interval_sec=0, gemini_api_key="k", gemini_fallback_models=(),
                 data_dir=main.settings.data_dir)
    state = {"fail": "quota"}

    def gemini_handler(request):
        if state["fail"] == "quota":
            return httpx.Response(429, json={"error": {"details": [{"violations": [{"quotaId": "PerDay"}]}]}})
        return httpx.Response(500, json={})

    gclient = httpx.Client(transport=httpx.MockTransport(gemini_handler))
    llm.fail = {"error_code": "startup_failed", "error": "x", "retryable": False}
    monkeypatch.setattr(pipeline, "get_provider", lambda: RoutingProvider(
        circuit=CircuitBreaker(tmp_path / "c.json"), gemini_factory=lambda: gemini.GeminiProvider(client=gclient),
        modal_factory=llm.provider))
    job_id = _submit(client)
    job = _wait(client, job_id)
    assert job["status"] == "failed" and job["error_code"] == "all_providers_failed"
    assert "hết hạn mức" in job["error"]
    d = jobs.job_dir(job_id)
    assert (d / "transcript.json").exists() and (d / "input.mp4").exists()  # completed work kept
    assert media.transcribe_calls == 1

    llm.fail = None
    assert client.post(f"/api/jobs/{job_id}/retry").status_code == 202
    job = _wait(client, job_id)
    assert job["status"] == "done", job["error"]
    assert media.transcribe_calls == 1 and media.preprocess_calls == 1  # transcript reused
    assert job["result"]["llm_route"]["fallback_reason"] in ("sticky", "circuit_open")


@needs_ffmpeg
def test_missing_cuda_fails_explicitly_and_keeps_the_video(cloud):
    client, media, _ = cloud
    media.transcribe_error = {"error_code": "cuda_unavailable", "error": "Không có CUDA.", "retryable": False}
    job_id = _submit(client)
    job = _wait(client, job_id)
    assert job["status"] == "failed" and job["error_code"] == "cuda_unavailable"
    assert "GPU/CUDA" in job["error"]
    d = jobs.job_dir(job_id)
    assert (d / "input.mp4").exists() and media.deleted  # remote temp files cleaned, source kept for retry


@needs_ffmpeg
def test_cpu_transcription_result_is_rejected(cloud):
    client, media, _ = cloud
    media.device = "cpu"
    job = _wait(client, _submit(client))
    assert job["status"] == "failed" and job["error_code"] == "cuda_unavailable"


@needs_ffmpeg
def test_cancel_during_remote_transcription_cancels_and_cleans_up(cloud):
    client, media, llm = cloud
    media.hold_transcription = True
    job_id = _submit(client)
    end = time.time() + 30
    while time.time() < end and media.transcribe_calls == 0:
        time.sleep(0.05)
    assert client.post(f"/api/jobs/{job_id}/cancel").status_code == 200
    job = _wait(client, job_id)
    assert job["status"] == "cancelled"
    assert media.cancelled and media.deleted and llm.submitted == []


@needs_ffmpeg
def test_transcript_is_redone_when_transcription_settings_change(cloud, monkeypatch):
    client, media, _ = cloud
    job_id = _submit(client)
    job = _wait(client, job_id)
    assert job["status"] == "done"
    d = jobs.job_dir(job_id)
    assert pipeline._transcript_current(d, pipeline._transcript_settings({"language": "vi"}, True))
    assert not pipeline._transcript_current(d, {**pipeline._transcript_settings({"language": "vi"}, True),
                                                "model": "medium"})


def test_invalid_upload_is_rejected_without_ffmpeg(cloud):
    client, media, _ = cloud
    resp = client.post("/api/jobs", files={"file": ("x.mp4", b"not a video at all", "video/mp4")})
    assert resp.status_code == 400
    assert media.uploads == {}


@needs_ffmpeg
def test_llm_gpu_is_prewarmed_during_transcription_only_when_it_will_be_used(cloud, monkeypatch):
    from app.ai import modal_ollama

    client, media, _ = cloud
    warmed = []
    monkeypatch.setattr(modal_ollama.ModalOllamaClient, "warmup", lambda self: warmed.append(1) or True)
    base = dict(media_processing="modal", modal_media_worker_url="https://media.modal.test",
                modal_poll_interval_sec=0, modal_llm_worker_url="https://llm.modal.test",
                modal_proxy_token_id="wk-test", modal_proxy_token_secret="ws-test", ai_provider="auto")
    use_settings(monkeypatch, gemini_api_key="", **base)  # no key: Modal will certainly be used
    job = _wait(client, _submit(client))
    assert job["status"] == "done" and job["result"]["media_processing"]["llm_prewarm"] is True
    time.sleep(0.2)
    assert warmed == [1]

    use_settings(monkeypatch, gemini_api_key="k", **base)  # Gemini first: no GPU is started in advance
    monkeypatch.setattr(pipeline, "get_provider", lambda: (_ for _ in ()).throw(AssertionError("stop")))
    _wait(client, _submit(client, title="khác"))
    assert warmed == [1]


def test_english_videos_use_the_english_whisper_model(monkeypatch):
    use_settings(monkeypatch, modal_whisper_model="small", modal_whisper_model_en="large-v3-turbo")
    assert pipeline._transcript_settings({"language": "en"}, True)["model"] == "large-v3-turbo"
    assert pipeline._transcript_settings({"language": "vi"}, True)["model"] == "small"
    assert pipeline._transcript_settings({"language": "auto"}, True)["model"] == "small"
