"""Low-level drawing primitives on top of python-pptx (all output stays editable)."""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt


def shape(slide, x: float, y: float, w: float, h: float, fill: RGBColor | None,
          kind=MSO_SHAPE.RECTANGLE, line: RGBColor | None = None, line_pt: float = 1.0,
          radius: float | None = None, alpha: float | None = None):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
        if alpha is not None:
            set_alpha(shp, alpha)
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_pt)
    shp.shadow.inherit = False
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    return shp


def set_alpha(shp, alpha: float) -> None:
    """Set fill opacity (0..1) on a solid-filled shape."""
    srgb = shp.fill._xPr.find(".//" + qn("a:srgbClr"))
    if srgb is None:
        return
    for old in srgb.findall(qn("a:alpha")):
        srgb.remove(old)
    el = etree.SubElement(srgb, qn("a:alpha"))
    el.set("val", str(int(alpha * 100000)))


def text(slide, x: float, y: float, w: float, h: float, paragraphs: list[str] | str, size: float,
         color: RGBColor, font: str, bold: bool = False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         bullet_color: RGBColor | None = None, bullet_char: str = "•", space_after: float = 0.45,
         italic: bool = False, target=None):
    """Add a text box (or fill ``target`` shape's frame) with consistent formatting."""
    if isinstance(paragraphs, str):
        paragraphs = [paragraphs]
    if target is None:
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        box = target
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1)
    tf.margin_top = tf.margin_bottom = Inches(0.05)

    indent = 0.32 if bullet_color is not None else 0.0
    for idx, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = 1.1
        if idx < len(paragraphs) - 1:
            p.space_after = Pt(size * space_after)
        run = p.add_run()
        run.text = para
        f = run.font
        f.size = Pt(size)
        f.bold = bold
        f.italic = italic
        f.name = font
        f.color.rgb = color
        if bullet_color is not None:
            _add_bullet(p, bullet_color, bullet_char, indent)
    return box


def _add_bullet(p, color: RGBColor, char: str, indent_in: float) -> None:
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Emu(Inches(indent_in)))))
    pPr.set("indent", str(-int(Emu(Inches(indent_in)))))
    bu_clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(bu_clr, qn("a:srgbClr")).set("val", str(color))
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", char)


def picture_cover(slide, path: Path, x: float, y: float, w: float, h: float):
    """Place an image filling the box exactly, cropping the overflow (like CSS cover)."""
    with Image.open(path) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    box_ratio, img_ratio = w / h, iw / ih
    if img_ratio > box_ratio:
        c = (1 - box_ratio / img_ratio) / 2
        pic.crop_left = pic.crop_right = c
    elif img_ratio < box_ratio:
        c = (1 - img_ratio / box_ratio) / 2
        pic.crop_top = pic.crop_bottom = c
    return pic


def image_ratio(path: Path) -> float:
    with Image.open(path) as im:
        return im.width / im.height


def set_background(slide, color: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color
