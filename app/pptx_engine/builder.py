"""Turn a validated ``Presentation`` outline into a designed, editable .pptx file."""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from pathlib import Path

from pptx import Presentation as PptxPresentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from ..ai.schema import Presentation, Slide
from . import draw
from .textfit import fit_or_trim, text_height_in
from .themes import Theme, get_theme

W, H = 13.333, 7.5
M = 0.6               # outer margin
CT, CB = 1.8, 6.85    # content top / bottom
CH = CB - CT

LABELS = {
    "vi": {"frame": "Khung hình tại", "sources": "Thời điểm trong video"},
    "en": {"frame": "Frame at", "sources": "Video timestamps"},
}


@dataclass
class FrameRef:
    path: Path
    timestamp: float


@dataclass
class BuildContext:
    theme: Theme
    deck_title: str
    language: str
    frames: list[FrameRef]
    total: int
    used_frames: set[Path] = field(default_factory=set)
    warnings: list[str] = field(default_factory=list)
    section_no: int = 0

    @property
    def labels(self) -> dict[str, str]:
        return LABELS.get(self.language, LABELS["en"])


def _fmt_ts(sec: float) -> str:
    m, s = divmod(int(sec), 60)
    return f"{m:02d}:{s:02d}"


# --------------------------------------------------------------------------- frames

def pick_frame(ctx: BuildContext, s: Slide, max_distance: float | None = None) -> FrameRef | None:
    """Unused frame nearest to the slide's source timestamps."""
    candidates = [f for f in ctx.frames if f.path not in ctx.used_frames]
    if not candidates:
        return None
    if s.source_timestamps:
        def dist(f: FrameRef) -> float:
            return min(abs(f.timestamp - t) for t in s.source_timestamps)
        best = min(candidates, key=dist)
        if max_distance is not None and dist(best) > max_distance:
            return None
    elif max_distance is not None:
        return None
    else:
        best = candidates[0]
    ctx.used_frames.add(best.path)
    return best


def plan_layouts(pres: Presentation, frames: list[FrameRef]) -> None:
    """Upgrade plain bullet slides to image layouts when a matching frame exists,
    so the deck alternates visual rhythm instead of repeating one layout."""
    if not frames:
        return
    reserved = sum(1 for s in pres.slides if s.layout_type in ("image_left", "image_right"))
    spare = len(frames) - reserved - 1  # keep one for the title slide
    side = "image_right"
    prev = None
    for s in pres.slides:
        if spare <= 0:
            break
        if s.layout_type == "bullets" and s.source_timestamps and 1 <= len(s.key_points) <= 5 and prev == "bullets":
            if any(min(abs(f.timestamp - t) for t in s.source_timestamps) <= 45 for f in frames):
                s.layout_type = side
                s.visual_type = "video_frame"
                side = "image_left" if side == "image_right" else "image_right"
                spare -= 1
        prev = s.layout_type


# --------------------------------------------------------------------------- chrome

def _title_color(t: Theme):
    return t.rgb("text") if t.dark else t.rgb("primary")


def _fit(ctx: BuildContext, slide_no: int, paras: list[str], w: float, h: float, max_pt: float,
         min_pt: float, bold: bool = False, indent: float = 0.0) -> tuple[list[str], float]:
    fitted, size = fit_or_trim(paras, w, h, max_pt, min_pt, bold, indent)
    if fitted != [p for p in paras if p]:
        ctx.warnings.append(f"Slide {slide_no}: text shortened to fit")
    return fitted, size


def _take(ctx: BuildContext, s: Slide, seq: list, limit: int) -> list:
    """Cap list length for a layout, recording a warning when content is dropped."""
    if len(seq) > limit:
        ctx.warnings.append(f"Slide {s.slide_number}: text shortened to fit ({len(seq) - limit} items dropped)")
    return seq[:limit]


def chrome(ctx: BuildContext, sl, s: Slide, title: str | None = None) -> None:
    """Background, title, accent and footer shared by content slides."""
    t = ctx.theme
    draw.set_background(sl, t.rgb("background"))
    if t.key == "professional":
        draw.shape(sl, 0, 0, W, 0.09, t.rgb("primary"))
    elif t.key == "creative":
        draw.shape(sl, W - 2.6, -1.6, 4.2, 4.2, t.rgb("accent2"), MSO_SHAPE.OVAL, alpha=0.18)
        draw.shape(sl, W - 1.2, 0.9, 1.4, 1.4, t.rgb("accent"), MSO_SHAPE.OVAL, alpha=0.22)

    title_w = W - 2 * M - (1.6 if t.key == "creative" else 0)
    paras, size = _fit(ctx, s.slide_number, [title or s.title], title_w, 1.0, 32, 22, bold=True)
    draw.text(sl, M, 0.45, title_w, 1.0, paras, size, _title_color(t), t.title_font, bold=True,
              anchor=MSO_ANCHOR.BOTTOM)
    if t.key == "academic":
        draw.shape(sl, M + 0.1, 1.52, W - 2 * M - 0.2, 0.015, t.rgb("accent"))
        draw.shape(sl, M + 0.1, 1.49, 1.1, 0.07, t.rgb("primary"))
    else:
        draw.shape(sl, M + 0.1, 1.5, 0.9, 0.07, t.rgb("accent"),
                   MSO_SHAPE.ROUNDED_RECTANGLE if t.rounded else MSO_SHAPE.RECTANGLE, radius=0.5)

    # Footer
    if t.key == "academic":
        draw.shape(sl, M + 0.1, 7.0, W - 2 * M - 0.2, 0.01, t.rgb("accent"))
    draw.text(sl, M, 7.03, 8, 0.35, ctx.deck_title[:90], 10, t.rgb("muted"), t.body_font)
    draw.text(sl, W - M - 1.5, 7.03, 1.5, 0.35, f"{s.slide_number} / {ctx.total}", 10, t.rgb("muted"),
              t.body_font, align=PP_ALIGN.RIGHT)


def _card(ctx: BuildContext, sl, x, y, w, h, fill_name: str = "surface", line: bool = False):
    t = ctx.theme
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if t.rounded else MSO_SHAPE.RECTANGLE
    radius = min(0.12, 0.25 / max(min(w, h), 0.5)) if t.rounded else None
    return draw.shape(sl, x, y, w, h, t.rgb(fill_name), kind, radius=radius,
                      line=t.rgb("accent") if line else None, line_pt=0.75)


def _badge(ctx: BuildContext, sl, x, y, d, label: str, color_name: str = "accent"):
    t = ctx.theme
    circ = draw.shape(sl, x, y, d, d, t.rgb(color_name), MSO_SHAPE.OVAL)
    draw.text(sl, 0, 0, 0, 0, label, max(11, d * 30), t.rgb("on_primary"), t.body_font, bold=True,
              align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, target=circ)
    circ.text_frame.margin_left = circ.text_frame.margin_right = 0
    return circ


def _bullets(ctx: BuildContext, sl, s: Slide, points: list[str], x, y, w, h, max_pt=24, min_pt=14,
             color_name="text", anchor=MSO_ANCHOR.TOP):
    t = ctx.theme
    paras, size = _fit(ctx, s.slide_number, points, w, h, max_pt, min_pt, indent=0.32)
    draw.text(sl, x, y, w, h, paras, size, t.rgb(color_name), t.body_font, anchor=anchor,
              bullet_color=t.rgb("accent"))


# --------------------------------------------------------------------------- layouts

def layout_title(ctx: BuildContext, sl, s: Slide) -> None:
    t = ctx.theme
    draw.set_background(sl, t.rgb("primary"))
    frame = pick_frame(ctx, s)
    if frame:
        # Show the whole frame (no heavy crop) as a framed 16:9 picture.
        img_w = 5.9
        img_h = img_w * 9 / 16
        ix, iy = W - M - img_w, (H - img_h) / 2
        draw.shape(sl, W - 3.4, -1.0, 4.6, 4.6, t.rgb("accent2"), MSO_SHAPE.OVAL, alpha=0.2)
        draw.shape(sl, ix + 0.22, iy + 0.22, img_w, img_h, t.rgb("accent"))
        draw.picture_cover(sl, frame.path, ix, iy, img_w, img_h)
        text_w = ix - M - 0.7
    else:
        draw.shape(sl, W - 4.8, -1.2, 6.0, 6.0, t.rgb("accent"), MSO_SHAPE.OVAL, alpha=0.25)
        draw.shape(sl, W - 3.0, 3.9, 4.2, 4.2, t.rgb("accent2"), MSO_SHAPE.OVAL, alpha=0.3)
        text_w = W - 2 * M - 3.5

    x = M + 0.2
    title = s.title or ctx.deck_title
    paras, size = _fit(ctx, s.slide_number, [title], text_w, 2.9, 44, 26, bold=True)
    draw.text(sl, x, 1.2, text_w, 2.9, paras, size, t.rgb("on_primary"), t.title_font, bold=True,
              anchor=MSO_ANCHOR.BOTTOM)
    draw.shape(sl, x + 0.1, 4.3, 1.3, 0.08, t.rgb("accent"))
    subtitle = s.subtitle or (s.key_points[0] if s.key_points else "")
    if subtitle:
        paras, size = _fit(ctx, s.slide_number, [subtitle], text_w, 1.6, 22, 14)
        draw.text(sl, x, 4.55, text_w, 1.6, paras, size, t.rgb("on_primary"), t.body_font)


def layout_section(ctx: BuildContext, sl, s: Slide) -> None:
    t = ctx.theme
    ctx.section_no += 1
    draw.set_background(sl, t.rgb("primary"))
    draw.shape(sl, W - 5.0, 1.0, 6.5, 6.5, t.rgb("accent2"), MSO_SHAPE.OVAL, alpha=0.18)
    draw.text(sl, M + 0.2, 1.3, 4, 1.4, f"{ctx.section_no:02d}", 72, t.rgb("accent"), t.title_font, bold=True,
              anchor=MSO_ANCHOR.BOTTOM)
    paras, size = _fit(ctx, s.slide_number, [s.title], 9.5, 1.9, 42, 26, bold=True)
    draw.text(sl, M + 0.2, 2.85, 9.5, 1.9, paras, size, t.rgb("on_primary"), t.title_font, bold=True)
    sub = s.subtitle or " · ".join(s.key_points[:3])
    if sub:
        paras, size = _fit(ctx, s.slide_number, [sub], 9.5, 1.3, 20, 13)
        draw.text(sl, M + 0.2, 4.9, 9.5, 1.3, paras, size, t.rgb("on_primary"), t.body_font)


def layout_bullets(ctx: BuildContext, sl, s: Slide) -> None:
    t = ctx.theme
    chrome(ctx, sl, s)
    points = _take(ctx, s, s.key_points, 6)
    if 2 <= len(points) <= 4:
        # Numbered rows read better than a bare list for a few key points.
        gap = 0.22
        row_h = min(1.15, (CH - gap * (len(points) - 1)) / len(points))
        y = CT + (CH - (row_h * len(points) + gap * (len(points) - 1))) / 2
        for i, p in enumerate(points):
            _card(ctx, sl, M + 0.1, y, W - 2 * M - 0.2, row_h)
            draw.shape(sl, M + 0.1, y, 0.09, row_h, t.rgb("accent" if i % 2 == 0 else "accent2"))
            d = min(0.6, row_h - 0.3)
            _badge(ctx, sl, M + 0.45, y + (row_h - d) / 2, d, str(i + 1), "accent" if i % 2 == 0 else "accent2")
            tw = W - 2 * M - 1.5
            paras, size = _fit(ctx, s.slide_number, [p], tw, row_h - 0.1, 22, 13)
            draw.text(sl, M + 1.3, y + 0.05, tw, row_h - 0.1, paras, size, t.rgb("text"), t.body_font,
                      anchor=MSO_ANCHOR.MIDDLE)
            y += row_h + gap
    else:
        _bullets(ctx, sl, s, points, M + 0.1, CT + 0.1, W - 2 * M - 0.2, CH - 0.2, max_pt=24, min_pt=14)


def layout_image(ctx: BuildContext, sl, s: Slide, image_side: str) -> None:
    t = ctx.theme
    frame = pick_frame(ctx, s)
    if frame is None:
        return layout_bullets(ctx, sl, s)
    chrome(ctx, sl, s)
    img_w = 6.3
    ratio = draw.image_ratio(frame.path)
    img_h = min(img_w / ratio, CH - 0.6)
    img_w = img_h * ratio if img_h == CH - 0.6 else img_w
    iy = CT + (CH - 0.45 - img_h) / 2
    if image_side == "left":
        ix, tx = M + 0.1, M + 0.1 + img_w + 0.5
    else:
        ix, tx = W - M - 0.1 - img_w, M + 0.1
    tw = W - 2 * M - 0.2 - img_w - 0.5

    draw.shape(sl, ix + 0.18, iy + 0.18, img_w, img_h, t.rgb("accent"), alpha=0.35)
    pic = draw.picture_cover(sl, frame.path, ix, iy, img_w, img_h)
    pic.line.color.rgb = t.rgb("surface")
    pic.line.width = Pt(1)
    draw.text(sl, ix, iy + img_h + 0.22, img_w, 0.35, f"{ctx.labels['frame']} {_fmt_ts(frame.timestamp)}",
              11, t.rgb("muted"), t.body_font, italic=True)
    _bullets(ctx, sl, s, _take(ctx, s, s.key_points, 5), tx, CT + 0.1, tw, CH - 0.3, max_pt=22, min_pt=13,
             anchor=MSO_ANCHOR.MIDDLE)


def layout_two_column(ctx: BuildContext, sl, s: Slide) -> None:
    t = ctx.theme
    chrome(ctx, sl, s)
    gap = 0.5
    cw = (W - 2 * M - 0.2 - gap) / 2
    need = max(text_height_in(p[:6], cw - 0.5, 22, indent_in=0.32) for p in (s.left_points, s.right_points))
    card_h = min(CH - 0.1, max(3.2, 1.5 + need))
    top = CT + (CH - card_h) / 2
    for i, (head, pts, color) in enumerate(
        [(s.left_title, s.left_points, "primary"), (s.right_title, s.right_points, "accent")]
    ):
        x = M + 0.1 + i * (cw + gap)
        _card(ctx, sl, x, top, cw, card_h)
        band = draw.shape(sl, x, top, cw, 0.8, t.rgb(color),
                          MSO_SHAPE.ROUNDED_RECTANGLE if t.rounded else MSO_SHAPE.RECTANGLE,
                          radius=0.2 if t.rounded else None)
        paras, size = _fit(ctx, s.slide_number, [head or "—"], cw - 0.3, 0.7, 22, 14, bold=True)
        draw.text(sl, 0, 0, 0, 0, paras, size, t.rgb("on_primary"), t.title_font, bold=True,
                  anchor=MSO_ANCHOR.MIDDLE, target=band)
        _bullets(ctx, sl, s, _take(ctx, s, pts, 6), x + 0.25, top + 1.05, cw - 0.5, card_h - 1.25, max_pt=22, min_pt=12)


def layout_process(ctx: BuildContext, sl, s: Slide) -> None:
    t = ctx.theme
    chrome(ctx, sl, s)
    items = s.items[:6]
    n = len(items)
    gap = 0.45
    bw = (W - 2 * M - 0.2 - gap * (n - 1)) / n
    d = 0.62
    need = max(text_height_in([it.text], bw - 0.2, 17) if it.text else 0 for it in items)
    proc_h = min(CB - (CT + 0.1 + d + 0.3), max(2.6, 1.6 + need))
    cy = CT + (CH - (d + 0.3 + proc_h)) / 2
    # Connector line behind the step badges.
    line = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(M + 0.1 + bw / 2), Inches(cy + d / 2),
                                   Inches(W - M - 0.1 - bw / 2), Inches(cy + d / 2))
    line.line.color.rgb = t.rgb("muted")
    line.line.width = Pt(1.5)
    for i, it in enumerate(items):
        x = M + 0.1 + i * (bw + gap)
        color = "accent" if i % 2 == 0 else "accent2"
        _badge(ctx, sl, x + (bw - d) / 2, cy, d, str(i + 1), color)
        card_y = cy + d + 0.3
        card_h = proc_h
        _card(ctx, sl, x, card_y, bw, card_h)
        draw.shape(sl, x, card_y, bw, 0.08, t.rgb(color))
        hp, hs = _fit(ctx, s.slide_number, [it.heading], bw - 0.2, 1.1, 20, 12, bold=True)
        draw.text(sl, x + 0.1, card_y + 0.2, bw - 0.2, 1.1, hp, hs, _title_color(t), t.title_font, bold=True)
        if it.text:
            tp, ts = _fit(ctx, s.slide_number, [it.text], bw - 0.2, card_h - 1.5, 18, 11)
            draw.text(sl, x + 0.1, card_y + 1.35, bw - 0.2, card_h - 1.5, tp, ts, t.rgb("text"), t.body_font)
        if i < n - 1:
            draw.shape(sl, x + bw + 0.08, card_y + card_h / 2 - 0.18, gap - 0.16, 0.36, t.rgb(color),
                       MSO_SHAPE.RIGHT_ARROW)


def layout_timeline(ctx: BuildContext, sl, s: Slide) -> None:
    t = ctx.theme
    chrome(ctx, sl, s)
    items = s.items[:6]
    n = len(items)
    slot = (W - 2 * M - 0.2) / n
    y_line = CT + CH * 0.45
    draw.shape(sl, M + 0.1, y_line - 0.03, W - 2 * M - 0.2, 0.06, t.rgb("muted"), alpha=0.6)
    for i, it in enumerate(items):
        cx = M + 0.1 + (i + 0.5) * slot
        color = "accent" if i % 2 == 0 else "accent2"
        draw.shape(sl, cx - 0.2, y_line - 0.2, 0.4, 0.4, t.rgb(color), MSO_SHAPE.OVAL)
        draw.shape(sl, cx - 0.08, y_line - 0.08, 0.16, 0.16, t.rgb("background"), MSO_SHAPE.OVAL)
        hp, hs = _fit(ctx, s.slide_number, [it.heading], slot - 0.2, y_line - CT - 0.45, 24, 12, bold=True)
        draw.text(sl, cx - slot / 2 + 0.1, CT + 0.15, slot - 0.2, y_line - CT - 0.45, hp, hs, t.rgb(color), t.title_font,
                  bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        if it.text:
            top = y_line + 0.45
            tp, ts = _fit(ctx, s.slide_number, [it.text], slot - 0.2, CB - top, 18, 11)
            draw.text(sl, cx - slot / 2 + 0.1, top, slot - 0.2, CB - top, tp, ts, t.rgb("text"), t.body_font,
                      align=PP_ALIGN.CENTER)


def _grid(n: int) -> tuple[int, int]:
    cols = n if n <= 3 else (2 if n == 4 else 3)
    return cols, math.ceil(n / cols)


def layout_cards(ctx: BuildContext, sl, s: Slide) -> None:
    t = ctx.theme
    chrome(ctx, sl, s)
    items = s.items[:6]
    cols, rows = _grid(len(items))
    gap = 0.35
    cw = (W - 2 * M - 0.2 - gap * (cols - 1)) / cols
    max_ch = (CH - gap * (rows - 1)) / rows
    # Shrink cards to their content so short texts don't leave empty boxes.
    need = max(text_height_in([it.text], cw - 0.5, 17) if it.text else 0 for it in items)
    ch = min(max_ch, max(1.9, 1.45 + need))
    top = CT + (CH - (ch * rows + gap * (rows - 1))) / 2
    for i, it in enumerate(items):
        r, c = divmod(i, cols)
        x = M + 0.1 + c * (cw + gap)
        y = top + r * (ch + gap)
        color = "accent" if (r + c) % 2 == 0 else "accent2"
        _card(ctx, sl, x, y, cw, ch)
        draw.shape(sl, x, y, cw, 0.08, t.rgb(color))
        d = 0.55
        _badge(ctx, sl, x + 0.25, y + 0.3, d, str(i + 1), color)
        hp, hs = _fit(ctx, s.slide_number, [it.heading], cw - 1.15, 0.85, 20, 12, bold=True)
        draw.text(sl, x + 0.95, y + 0.25, cw - 1.15, 0.85, hp, hs, _title_color(t), t.title_font, bold=True,
                  anchor=MSO_ANCHOR.MIDDLE)
        if it.text:
            body_h = ch - 1.35
            tp, ts = _fit(ctx, s.slide_number, [it.text], cw - 0.5, body_h, 18, 11)
            draw.text(sl, x + 0.25, y + 1.2, cw - 0.5, body_h, tp, ts, t.rgb("text"), t.body_font)


def layout_stats(ctx: BuildContext, sl, s: Slide) -> None:
    t = ctx.theme
    chrome(ctx, sl, s)
    items = s.items[:6]
    cols, rows = _grid(len(items)) if len(items) != 4 else (4, 1)
    gap = 0.35
    cw = (W - 2 * M - 0.2 - gap * (cols - 1)) / cols
    ch = min(4.2, (CH - gap * (rows - 1)) / rows)
    top = CT + (CH - (ch * rows + gap * (rows - 1))) / 2
    for i, it in enumerate(items):
        r, c = divmod(i, cols)
        x = M + 0.1 + c * (cw + gap)
        y = top + r * (ch + gap)
        color = "accent" if i % 2 == 0 else "accent2"
        _card(ctx, sl, x, y, cw, ch)
        vh = ch * 0.45
        vp, vs = _fit(ctx, s.slide_number, [it.heading], cw - 0.3, vh, 54, 24, bold=True)
        draw.text(sl, x + 0.15, y + 0.2, cw - 0.3, vh, vp, vs, t.rgb(color), t.title_font, bold=True,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        draw.shape(sl, x + cw / 2 - 0.4, y + 0.3 + vh, 0.8, 0.05, t.rgb(color))
        if it.text:
            lh = ch - vh - 0.7
            lp, ls = _fit(ctx, s.slide_number, [it.text], cw - 0.4, lh, 18, 11)
            draw.text(sl, x + 0.2, y + vh + 0.5, cw - 0.4, lh, lp, ls, t.rgb("text"), t.body_font,
                      align=PP_ALIGN.CENTER)


def layout_summary(ctx: BuildContext, sl, s: Slide) -> None:
    t = ctx.theme
    chrome(ctx, sl, s)
    points = _take(ctx, s, s.key_points, 6)
    two_cols = len(points) > 3
    cols = 2 if two_cols else 1
    rows = math.ceil(len(points) / cols)
    gap = 0.3
    cw = (W - 2 * M - 0.2 - gap * (cols - 1)) / cols
    rh = min(1.3, (CH - gap * (rows - 1)) / rows)
    top = CT + (CH - (rh * rows + gap * (rows - 1))) / 2
    for i, p in enumerate(points):
        c, r = (i // rows, i % rows) if two_cols else (0, i)
        x = M + 0.1 + c * (cw + gap)
        y = top + r * (rh + gap)
        _card(ctx, sl, x, y, cw, rh)
        d = min(0.55, rh - 0.3)
        _badge(ctx, sl, x + 0.25, y + (rh - d) / 2, d, "✓", "accent2")
        tw = cw - 1.1
        pp, ps = _fit(ctx, s.slide_number, [p], tw, rh - 0.1, 20, 12)
        draw.text(sl, x + 0.95, y + 0.05, tw, rh - 0.1, pp, ps, t.rgb("text"), t.body_font,
                  anchor=MSO_ANCHOR.MIDDLE)


LAYOUTS = {
    "title": layout_title,
    "section": layout_section,
    "bullets": layout_bullets,
    "image_left": lambda c, sl, s: layout_image(c, sl, s, "left"),
    "image_right": lambda c, sl, s: layout_image(c, sl, s, "right"),
    "two_column": layout_two_column,
    "process": layout_process,
    "timeline": layout_timeline,
    "cards": layout_cards,
    "stats": layout_stats,
    "summary": layout_summary,
}


# --------------------------------------------------------------------------- entry point

def build_pptx(pres: Presentation, out_path: Path, style: str = "professional",
               frames: list[FrameRef] | None = None) -> list[str]:
    """Render the deck to ``out_path``. Returns a list of non-fatal layout warnings."""
    frames = sorted(frames or [], key=lambda f: f.timestamp)
    theme = get_theme(style)
    plan_layouts(pres, frames)
    ctx = BuildContext(theme=theme, deck_title=pres.presentation_title, language=pres.language,
                       frames=frames, total=len(pres.slides))

    prs = PptxPresentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    prs.core_properties.title = pres.presentation_title
    prs.core_properties.subject = pres.subtitle or pres.presentation_title

    for s in pres.slides:
        sl = prs.slides.add_slide(blank)
        LAYOUTS.get(s.layout_type, layout_bullets)(ctx, sl, s)
        notes = s.speaker_notes.strip()
        if s.source_timestamps:
            stamps = ", ".join(_fmt_ts(x) for x in sorted(set(s.source_timestamps))[:6])
            notes = f"{notes}\n\n{ctx.labels['sources']}: {stamps}".strip()
        if notes:
            sl.notes_slide.notes_text_frame.text = notes

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return ctx.warnings
