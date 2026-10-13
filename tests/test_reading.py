"""Reading notes (schema, transcript fallback, notes text), full-width slides + notes in the PPTX, library."""

import dataclasses
import json
import time

import pytest
from pptx import Presentation as PptxPresentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

from app import jobs, library
from app.ai.schema import Presentation, TimeRange, gemini_response_schema
from app.pptx_engine import build_pptx, validate_pptx
from app.pptx_engine.builder import FULL_W, H
from app.reading import ensure_reading_notes, slide_payload, speaker_notes_text, transcript_excerpt
from app.transcribe import Segment, Transcript
from tests.sample_outline import SAMPLE, full_presentation, sample_presentation


def _transcript():
    return Transcript([Segment(i * 10.0, i * 10.0 + 10, f"Câu số {i} giải thích nội dung.") for i in range(30)], "vi", 300)


def test_outlines_without_reading_fields_still_load():
    p = sample_presentation()  # written before reading notes existed
    s = p.slides[1]
    assert s.brief_explanation == "" and s.reading_notes.is_empty() and s.source_ranges == []


def test_language_name_is_normalised():
    assert Presentation.model_validate({**SAMPLE, "language": "Vietnamese"}).language == "vi"
    assert Presentation.model_validate({**SAMPLE, "language": "en-US"}).language == "en"


def test_reading_fields_validate_and_trim():
    data = json.loads(json.dumps(SAMPLE))
    data["slides"][1].update({
        "brief_explanation": "  Một   câu ngắn. ",
        "reading_notes": {"summary": "Ý chính.", "detailed_explanation": "Đoạn 1.\n\n\nĐoạn   2.",
                          "examples": ["Ví dụ", " "] + ["x"] * 9, "key_takeaway": "Nhớ."},
        "source_ranges": [{"start": 30, "end": 12}],
    })
    s = Presentation.model_validate(data).slides[1]
    assert s.brief_explanation == "Một câu ngắn."
    assert s.reading_notes.detailed_explanation == "Đoạn 1.\n\nĐoạn 2."
    assert len(s.reading_notes.examples) == 6 and s.source_ranges[0].start == 12


def test_gemini_schema_requests_reading_notes():
    props = gemini_response_schema()["properties"]["slides"]["items"]
    assert {"brief_explanation", "reading_notes", "source_ranges"} <= set(props["properties"])
    assert "reading_notes" in props["required"]
    assert set(props["properties"]["reading_notes"]["properties"]) == {
        "summary", "detailed_explanation", "examples", "key_takeaway"}


def test_fallback_explanations_come_from_the_transcript_and_never_overwrite():
    p = sample_presentation()
    p.slides[3].reading_notes.detailed_explanation = "Viết bởi AI."
    ensure_reading_notes(p, _transcript())
    s = p.slides[1]  # source_timestamps [12.0]
    assert "Câu số 1" in s.reading_notes.detailed_explanation and s.brief_explanation
    assert p.slides[3].reading_notes.detailed_explanation == "Viết bởi AI."
    for slide in p.slides:  # every slide ends up with something to read
        assert slide.reading_notes.detailed_explanation or slide.reading_notes.summary


def test_excerpt_uses_source_ranges():
    p = sample_presentation()
    s = p.slides[1]
    s.source_ranges = [TimeRange(start=95, end=118)]
    text = transcript_excerpt(_transcript(), s)
    assert "Câu số 9" in text and "Câu số 11" in text and "Câu số 13" not in text


def test_speaker_notes_contain_full_narration_and_timestamps():
    s = sample_presentation().slides[1]
    s.reading_notes.detailed_explanation = "Đoạn giải thích."
    s.reading_notes.examples = ["Ví dụ A"]
    s.reading_notes.key_takeaway = "Điều cần nhớ."
    notes = speaker_notes_text(s, "vi")
    for part in ("Ba mục tiêu.", "Giải thích chi tiết:", "Đoạn giải thích.", "• Ví dụ A", "Ý chính cần nhớ: Điều cần nhớ.",
                 "Thời điểm trong video: 00:12"):
        assert part in notes
    payload = slide_payload(s, _transcript())
    assert payload["slide_id"] == "slide-02" and payload["reading_notes"]["examples"] == ["Ví dụ A"]
    assert payload["source_timestamps"] and payload["transcript_excerpt"]


def _explained(style="professional"):
    pres = full_presentation()
    for s in pres.slides:
        s.brief_explanation = f"Giải thích ngắn gọn cho slide “{s.title}” để người đọc hiểu ngay ý nghĩa."
        s.reading_notes.detailed_explanation = f"Giải thích chi tiết về {s.title}."
        s.reading_notes.key_takeaway = "Một điều cần nhớ ngắn gọn."
    return pres


SIDEBAR_LABELS = ("GIẢI THÍCH", "EXPLANATION")


def _illustrations(tmp_path, pres):
    from PIL import Image
    img = tmp_path / "illus.png"
    Image.new("RGB", (1344, 768), (90, 120, 160)).save(img)
    return {s.slide_id: img for s in pres.slides if s.layout_type in ("title", "concept", "image_left")}


@pytest.mark.parametrize("style", ["professional", "academic", "creative"])
@pytest.mark.parametrize("mode", ["reading", "presentation"])
def test_slides_use_full_canvas_without_explanation_sidebar(tmp_path, style, mode):
    """Regression: no explanation panel is drawn on the slide; explanations stay in the notes."""
    pres = _explained()
    out = tmp_path / f"{style}-{mode}.pptx"
    build_pptx(pres, out, style, [], illustrations=_illustrations(tmp_path, pres), layout_mode=mode)
    report = validate_pptx(out, expected_slides=len(pres.slides))
    assert report.ok, report.errors
    deck = PptxPresentation(str(out))
    assert round(deck.slide_width / deck.slide_height, 3) == round(16 / 9, 3)
    assert len({s.layout_type for s in pres.slides}) >= 5
    for idx, (sl, s) in enumerate(zip(deck.slides, pres.slides), start=1):
        texts = [sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        assert not any(sh.name.startswith("Reading panel") for sh in sl.shapes), f"slide {idx} has a sidebar"
        assert not any(t.strip() in SIDEBAR_LABELS for t in texts), f"slide {idx} has an explanation label"
        joined = "\n".join(texts)
        for note_only in ("Giải thích ngắn gọn", "Giải thích chi tiết", "Một điều cần nhớ"):
            assert note_only not in joined, f"slide {idx}: notes text duplicated on the canvas"
        for sh in sl.shapes:  # text and pictures stay on the 16:9 canvas (decorative shapes may bleed)
            if (sh.has_text_frame and sh.text_frame.text.strip()) or sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                assert Emu(sh.left).inches >= -0.01 and Emu(sh.left + sh.width).inches <= FULL_W + 0.01
                assert Emu(sh.top).inches >= -0.01 and Emu(sh.top + sh.height).inches <= H + 0.01
        if s.layout_type not in ("title", "section", "conclusion"):
            right = max(Emu(sh.left + sh.width).inches for sh in sl.shapes
                        if sh.has_text_frame and sh.text_frame.text.strip())
            assert right > 10.0, f"slide {idx} ({s.layout_type}) leaves the right side of the canvas empty"
        notes = sl.notes_slide.notes_text_frame.text
        assert f"Giải thích chi tiết về {s.title}" in notes
        assert "Ý chính cần nhớ: Một điều cần nhớ ngắn gọn." in notes
        if s.source_timestamps or s.source_ranges:
            assert "Thời điểm trong video:" in notes
        assert notes == speaker_notes_text(s, pres.language)


def test_library_roundtrip_and_finished_jobs_are_kept(monkeypatch):
    library.save("f" * 32, title="Bài A", language="vi", style="professional", layout_mode="reading",
                 slide_count=3, outline_json=sample_presentation().model_dump_json(), transcript_json=None,
                 source={"type": "upload"})
    library.set_google("f" * 32, {"presentation_id": "p1", "url": "https://docs.google.com/presentation/d/p1/edit"})
    entry = library.get("f" * 32)
    assert entry["title"] == "Bài A" and entry["google"]["presentation_id"] == "p1"
    assert any(i["job_id"] == "f" * 32 and i["google_url"] for i in library.recent())

    job_id, _ = jobs.new_job({}, "x.mp4")
    jobs.update(job_id, status="done", created_at=time.time() - 10 * 86400)
    failed_id, _ = jobs.new_job({}, "y.mp4")
    jobs.update(failed_id, status="failed", created_at=time.time() - 10 * 86400)
    jobs.cleanup_expired()
    assert jobs.job_dir(job_id).exists() and not jobs.job_dir(failed_id).exists()
    monkeypatch.setattr(jobs, "settings", dataclasses.replace(jobs.settings, keep_finished_presentations=False))
    jobs.cleanup_expired()
    assert not jobs.job_dir(job_id).exists()
    library.delete("f" * 32)
    assert library.get("f" * 32) is None
