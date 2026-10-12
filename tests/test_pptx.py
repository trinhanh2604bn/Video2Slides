import pytest
from PIL import Image
from pptx import Presentation as Pptx

from app.ai.base import OutlineRequest
from app.ai.offline import OfflineProvider
from app.ai.schema import Presentation
from app.pptx_engine import THEMES, FrameRef, build_pptx, validate_pptx
from app.pptx_engine.textfit import fit_or_trim, wrapped_lines
from app.transcribe import Segment, Transcript
from tests.sample_outline import SAMPLE, sample_presentation


@pytest.fixture
def frames(tmp_path):
    out = []
    for i, (ts, color) in enumerate([(5, "navy"), (40, "orange"), (200, "green")]):
        p = tmp_path / f"f{i}.jpg"
        Image.new("RGB", (1280, 720), color).save(p)
        out.append(FrameRef(p, ts))
    return out


@pytest.mark.parametrize("style", list(THEMES))
def test_build_all_layouts_all_themes(tmp_path, frames, style):
    out = tmp_path / f"{style}.pptx"
    warnings = build_pptx(sample_presentation(), out, style, frames)
    report = validate_pptx(out, expected_slides=11)
    assert report.ok, report.errors
    assert report.pictures == 3
    assert not report.warnings, report.warnings
    assert not warnings


def test_output_is_editable_with_notes(tmp_path, frames):
    out = tmp_path / "d.pptx"
    build_pptx(sample_presentation(), out, "professional", frames)
    prs = Pptx(str(out))
    assert prs.slide_width / prs.slide_height == pytest.approx(16 / 9, rel=0.01)
    all_text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Quy trình xây dựng mô hình" in all_text  # real text, not pictures of text
    assert "Thời điểm trong video" in prs.slides[3].notes_slide.notes_text_frame.text


def test_image_layout_without_frames_falls_back(tmp_path):
    out = tmp_path / "noframes.pptx"
    build_pptx(sample_presentation(), out, "academic", [])
    report = validate_pptx(out, expected_slides=11)
    assert report.ok and report.pictures == 0


def test_long_text_is_trimmed_not_overflowing(tmp_path):
    long = "Đây là một câu rất dài được lặp lại nhiều lần để kiểm tra khả năng chống tràn chữ " * 3
    data = dict(SAMPLE)
    data["slides"] = [SAMPLE["slides"][0], {
        "slide_number": 2, "title": long[:119], "layout_type": "bullets",
        "key_points": [long[:239]] * 8, "speaker_notes": "", "visual_type": "none",
    }]
    out = tmp_path / "long.pptx"
    warnings = build_pptx(Presentation.model_validate(data), out)
    report = validate_pptx(out)
    assert report.ok
    assert not report.warnings
    assert any("shortened" in w for w in warnings)


def test_fit_or_trim():
    paras, size = fit_or_trim(["ngắn"], 5, 1, 24, 12)
    assert paras == ["ngắn"] and size == 24
    paras, size = fit_or_trim(["từ " * 400], 3, 0.5, 20, 12)
    assert size >= 12 and paras[0].endswith("…") and len(paras[0]) < 400 * 3


def test_wrapped_lines_counts_long_words():
    assert wrapped_lines("a" * 200, 1.2, 12) > 1
    assert wrapped_lines("word " * 50, 2, 12) > wrapped_lines("word " * 5, 2, 12)


def test_offline_provider_produces_valid_deck(tmp_path, frames):
    segs = [Segment(i * 10, i * 10 + 9, f"Đây là câu số {i} trong bài giảng về chủ đề thử nghiệm.") for i in range(30)]
    t = Transcript(segs, "vi", 300)
    pres = OfflineProvider().generate(OutlineRequest(t, "vi", None, frame_timestamps=(5, 40, 200)))
    assert pres.slides[0].layout_type == "title"
    assert pres.slides[-1].layout_type == "summary"
    out = tmp_path / "off.pptx"
    build_pptx(pres, out, "creative", frames)
    assert validate_pptx(out, expected_slides=len(pres.slides)).ok


def test_validate_rejects_non_pptx(tmp_path):
    bad = tmp_path / "bad.pptx"
    bad.write_bytes(b"nope")
    assert not validate_pptx(bad).ok
