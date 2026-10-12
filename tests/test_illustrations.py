"""AI illustration pipeline: planner, prompts, provider, cache, retries, limits, PPTX, API."""

import base64
import dataclasses
import io
import json
import time

import httpx
import pytest
from PIL import Image
from pptx import Presentation as Pptx

from app import config
from app.ai.prompt import build_user_prompt, illustration_rule
from app.ai.schema import Presentation, gemini_response_schema
from app.illustrations import usage
from app.illustrations.base import ImageGenerationError, ImageJobStatus, ImageRequest
from app.illustrations.mock import MockImageProvider
from app.illustrations.modal_sdxl import ModalSDXLImageProvider
from app.illustrations.planner import (
    FALLBACK_TITLE_SUBJECT, apply_illustrations, deck_seed, image_cap, plan_illustrations,
)
from app.illustrations.service import IllustrationService, cache_dir
from app.illustrations.validation import check_params, check_prompt, sanitize_image
from app.pptx_engine import build_pptx, validate_pptx
from tests.sample_outline import SAMPLE


@pytest.fixture(autouse=True)
def image_settings(monkeypatch, tmp_path):
    s = dataclasses.replace(
        config.settings, data_dir=tmp_path / "data", ai_images_enabled=True, image_provider="mock",
        image_poll_interval_sec=0, image_batch_timeout_sec=5, image_kill_switch=False,
        image_daily_limit=40, image_monthly_limit=300, image_monthly_budget_usd=5.0, image_max_attempts=2,
        image_on_failure="fail", modal_image_worker_url="https://ws--worker.modal.run",
        modal_proxy_token_id="wk-test", modal_proxy_token_secret="ws-test",
    )
    monkeypatch.setattr(config, "settings", s)
    return s


def outline_with_prompts() -> Presentation:
    data = json.loads(json.dumps(SAMPLE))
    data["slides"][0]["image_prompt"] = "a sunrise over a modern city of glowing data streams"
    data["slides"][1]["image_prompt"] = "a student looking at a bright path of stepping stones"  # bullets
    data["slides"][3]["image_prompt"] = "a robot learning to sort colourful shapes on a table"   # image_left
    data["slides"][3]["visual_type"] = "illustration"
    data["slides"][8]["image_prompt"] = "should be ignored: stats slide"                     # stats
    return Presentation.model_validate(data)


def jpeg(w=1344, h=768, color=None) -> bytes:
    im = Image.new("RGB", (w, h), color or "white")
    if color is None:
        for x in range(0, w, 32):
            im.paste((x % 255, 80, 160), (x, 0, x + 16, h))
    buf = io.BytesIO()
    im.save(buf, "JPEG")
    return buf.getvalue()


# ------------------------------------------------------------------ planner & prompts

def test_planner_selects_llm_prompts_and_skips_text_layouts():
    plan = plan_illustrations(outline_with_prompts(), "a" * 32, "professional", "high")
    ids = [r.slide_id for r in plan.requests]
    assert ids == ["slide-01", "slide-02", "slide-04"]  # stats slide excluded
    r = plan.requests[0]
    assert r.presentation_id == "a" * 32 and r.slide_title
    assert r.prompt.startswith("a sunrise over a modern city")
    assert "corporate illustration" in r.prompt          # shared deck style
    assert "text" in r.negative_prompt and "watermark" in r.negative_prompt
    assert (r.width, r.height, r.aspect_ratio, r.output_format) == (1344, 768, "16:9", "jpeg")
    assert all(req.prompt_source == "llm" for req in plan.requests)
    # Slide titles never go into the raster prompt.
    assert all(req.slide_title not in req.prompt for req in plan.requests)


def test_planner_caps_density_and_global_limit():
    pres = outline_with_prompts()
    assert len(plan_illustrations(pres, "x" * 32, "academic", "high", max_images=2).requests) == 2
    assert image_cap(11, "standard", 8) == 4
    assert image_cap(30, "high", 99) == 8      # hard cap of 8 per presentation
    assert image_cap(11, "standard", 0) == 0


def test_planner_fallback_without_llm_prompts():
    plan = plan_illustrations(Presentation.model_validate(SAMPLE), "b" * 32, "creative")
    assert [r.slide_id for r in plan.requests] == ["slide-01"]
    assert plan.requests[0].prompt_source == "fallback"
    assert plan.requests[0].prompt.startswith(FALLBACK_TITLE_SUBJECT)


def test_seeds_are_stable_and_variants_change_them():
    pres = outline_with_prompts()
    a = plan_illustrations(pres, "c" * 32, "professional")
    b = plan_illustrations(pres, "c" * 32, "professional")
    assert [r.seed for r in a.requests] == [r.seed for r in b.requests]
    assert [r.output_id for r in a.requests] == [r.output_id for r in b.requests]
    v = plan_illustrations(pres, "c" * 32, "professional", variants={"slide-01": 1})
    assert v.requests[0].seed != a.requests[0].seed == deck_seed("c" * 32, "slide-01")


def test_apply_illustrations_switches_layouts():
    pres = outline_with_prompts()
    apply_illustrations(pres, {"slide-02"})
    assert pres.slides[1].layout_type in ("image_left", "image_right")
    assert pres.slides[1].visual_type == "illustration"
    assert pres.slides[4].layout_type == "two_column"


def test_blocked_prompt_is_skipped_not_sent():
    data = json.loads(json.dumps(SAMPLE))
    data["slides"][0]["image_prompt"] = "a nude figure"
    plan = plan_illustrations(Presentation.model_validate(data), "d" * 32, "professional")
    assert not plan.requests and plan.skipped


def test_gemini_schema_and_prompt_only_change_when_requested():
    assert "image_prompt" not in json.dumps(gemini_response_schema())
    assert "image_prompt" in json.dumps(gemini_response_schema(include_image_prompt=True))
    assert illustration_rule(0) == ""
    rule = illustration_rule(4)
    assert "at most 4" in rule and "English" in rule and "Never ask for text" in rule


# ------------------------------------------------------------------ validation

def test_param_limits():
    check_params(1344, 768, 25, 6.0, 1, "jpeg")
    for bad in [(2048, 2048, 25, 6, 1, "jpeg"), (1340, 768, 25, 6, 1, "jpeg"), (1344, 768, 80, 6, 1, "jpeg"),
                (1344, 768, 25, 30, 1, "jpeg"), (1344, 768, 25, 6, -1, "jpeg"), (1344, 768, 25, 6, 1, "gif")]:
        with pytest.raises(ImageGenerationError):
            check_params(*bad)
    assert check_prompt("  a\x00 cat \n") == "a cat"
    with pytest.raises(ImageGenerationError):
        check_prompt("explicit gore scene")


def test_sanitize_image():
    data, w, h = sanitize_image(jpeg(), 1344, 768)
    assert (w, h) == (1344, 768) and data[:2] == b"\xff\xd8"
    with pytest.raises(ImageGenerationError, match="trống"):
        sanitize_image(jpeg(color="black"), 1344, 768)
    with pytest.raises(ImageGenerationError):
        sanitize_image(b"not an image" * 10, 1344, 768)
    with pytest.raises(ImageGenerationError):
        sanitize_image(jpeg(1024, 1024), 1344, 768)


# ------------------------------------------------------------------ Modal provider (HTTP contract)

def _modal(handler):
    return ModalSDXLImageProvider(client=httpx.Client(transport=httpx.MockTransport(handler)))


def _req(slide="slide-01", seed=1) -> ImageRequest:
    return ImageRequest("p" * 32, slide, "T", "", "a lighthouse", "text", "professional", "16:9", 1344, 768, seed,
                        25, 6.0)


def test_modal_provider_sends_proxy_auth_and_maps_states():
    seen = []

    def handler(request: httpx.Request):
        seen.append(request)
        assert request.headers["Modal-Key"] == "wk-test" and request.headers["Modal-Secret"] == "ws-test"
        if request.url.path == "/v1/generate":
            body = json.loads(request.content)
            assert body["request_id"] == "idem-1" and body["params"]["width"] == 1344
            assert "slide_title" not in body["params"]          # only pixel parameters leave the backend
            return httpx.Response(200, json={"call_id": "fc-abc123def456"})
        ref = request.url.path.rsplit("/", 1)[-1]
        if ref == "fc-pending0000":
            return httpx.Response(200, json={"status": "pending"})
        if ref == "fc-failed00000":
            return httpx.Response(200, json={"status": "failed", "error_code": "cuda_oom", "retryable": True})
        if ref == "fc-gone0000000":
            return httpx.Response(404, json={"detail": "unknown call"})
        return httpx.Response(200, json={"status": "done", "image_b64": base64.b64encode(jpeg()).decode(),
                                         "format": "jpeg", "width": 1344, "height": 768, "seed": 1,
                                         "model": "sdxl", "metrics": {"inference_seconds": 11.5}})

    p = _modal(handler)
    assert p.submit(_req(), idempotency_key="idem-1") == "fc-abc123def456"
    assert seen[0].url.host == "ws--worker.modal.run"
    assert p.get_job_status("fc-pending0000").state == "pending"
    st = p.get_job_status("fc-failed00000")
    assert st.state == "failed" and st.error_code == "cuda_oom" and st.retryable
    assert p.get_job_status("fc-gone0000000").state == "expired"
    done = p.get_job_status("fc-done0000000")
    assert done.state == "done" and done.result.metrics["inference_seconds"] == 11.5


def test_modal_provider_auth_and_transport_errors():
    p = _modal(lambda r: httpx.Response(401))
    with pytest.raises(ImageGenerationError) as e:
        p.submit(_req())
    assert e.value.code == "auth" and not e.value.retryable

    def boom(r):
        raise httpx.ConnectTimeout("slow")
    with pytest.raises(ImageGenerationError) as e:
        _modal(boom).get_job_status("fc-x0000000000")
    assert e.value.retryable

    p = _modal(lambda r: httpx.Response(200, json={"status": "done", "image_b64": "!!notbase64!!"}))
    assert p.get_job_status("fc-x0000000000").state == "failed"


def test_modal_provider_requires_configuration(monkeypatch):
    monkeypatch.setattr(config, "settings", dataclasses.replace(config.settings, modal_proxy_token_secret=""))
    with pytest.raises(ImageGenerationError) as e:
        ModalSDXLImageProvider()
    assert e.value.code == "not_configured"
    monkeypatch.setattr(config, "settings", dataclasses.replace(config.settings, modal_proxy_token_secret="x",
                                                                modal_image_worker_url="http://insecure"))
    with pytest.raises(ImageGenerationError):
        ModalSDXLImageProvider()


# ------------------------------------------------------------------ orchestration, cache, retries

def _plan(n=3):
    pres = outline_with_prompts()
    plan = plan_illustrations(pres, "e" * 32, "professional", "high")
    plan.requests = plan.requests[:n]
    return plan


def test_batch_generates_and_is_idempotent(tmp_path):
    provider = MockImageProvider()
    svc = IllustrationService(tmp_path / "job1", provider, sleep=lambda _: None)
    out = svc.run(_plan())
    assert sorted(out.images) == ["slide-01", "slide-02", "slide-04"] and not out.failed
    assert len(provider.submitted) == 3 and not out.real_model
    # Same job again (retry / refresh / restart): nothing is submitted twice.
    out2 = IllustrationService(tmp_path / "job1", provider, sleep=lambda _: None).run(_plan())
    assert len(provider.submitted) == 3 and len(out2.images) == 3
    # Another job with identical requests reuses the global cache.
    out3 = IllustrationService(tmp_path / "job2", provider, sleep=lambda _: None).run(_plan())
    assert len(provider.submitted) == 3 and out3.from_cache == 3
    assert len(list(cache_dir().glob("*.jpg"))) == 3


def test_resume_polling_after_restart_without_resubmitting(tmp_path):
    provider = MockImageProvider(pending_polls=1000)
    s = config.settings
    config.settings = dataclasses.replace(s, image_batch_timeout_sec=0)
    try:
        out = IllustrationService(tmp_path / "j", provider, sleep=lambda _: None).run(_plan(1))
    finally:
        config.settings = s
    assert out.failed and out.failed[0]["code"] == "timeout"
    manifest = json.loads((tmp_path / "j" / "illustrations" / "manifest.json").read_text(encoding="utf-8"))
    entry = manifest["slide-01"]
    # Simulate a backend restart while the GPU call was still running.
    entry["state"] = "submitted"
    (tmp_path / "j" / "illustrations" / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    provider._jobs[entry["job_ref"]]["polls"] = 1
    out = IllustrationService(tmp_path / "j", provider, sleep=lambda _: None).run(_plan(1))
    assert "slide-01" in out.images and len(provider.submitted) == 1


def test_retryable_gpu_failure_is_retried_once(tmp_path):
    provider = MockImageProvider(fail_codes={"slide-01": ["cuda_oom"]})
    out = IllustrationService(tmp_path / "j", provider, sleep=lambda _: None).run(_plan(1))
    assert "slide-01" in out.images and len(provider.submitted) == 2


def test_failures_are_bounded_and_reported(tmp_path):
    provider = MockImageProvider(fail_codes={"slide-01": ["cuda_oom", "cuda_oom", "cuda_oom"],
                                             "slide-02": ["invalid_params"]})
    out = IllustrationService(tmp_path / "j", provider, sleep=lambda _: None).run(_plan(3))
    assert len(provider.submitted) == 4            # 2 attempts for slide-01, 1 each for the others
    assert {f["slide_id"]: f["code"] for f in out.failed} == {"slide-01": "cuda_oom", "slide-02": "invalid_params"}
    assert list(out.images) == ["slide-04"]


def test_gpu_timeout_cancels_remaining_work(tmp_path, monkeypatch):
    provider = MockImageProvider(pending_polls=10_000)
    cancelled = []
    monkeypatch.setattr(provider, "cancel", cancelled.append)
    monkeypatch.setattr(config, "settings", dataclasses.replace(config.settings, image_batch_timeout_sec=0))
    out = IllustrationService(tmp_path / "j", provider, sleep=lambda _: None).run(_plan(2))
    assert [f["code"] for f in out.failed] == ["timeout", "timeout"] and len(cancelled) == 2


# ------------------------------------------------------------------ cost controls

def test_limits_block_before_any_gpu_call(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "settings", dataclasses.replace(config.settings, image_daily_limit=2))
    provider = MockImageProvider()
    out = IllustrationService(tmp_path / "j", provider, sleep=lambda _: None).run(_plan(3))
    assert not provider.submitted and {f["code"] for f in out.failed} == {"limit"}


def test_budget_estimate_and_kill_switch():
    assert usage.estimate_cost_usd(0) == 0
    assert 0 < usage.estimate_cost_usd(8) < 0.1          # ~8 images on an L4 incl. cold start + idle window
    assert usage.availability().available
    usage.set_kill_switch(True)
    assert not usage.availability().available
    with pytest.raises(ImageGenerationError) as e:
        usage.reserve(1)
    assert e.value.code == "limit"
    usage.set_kill_switch(False)
    usage.reserve(3)
    assert usage.report()["today"]["images"] == 3


def test_monthly_budget(monkeypatch):
    monkeypatch.setattr(config, "settings", dataclasses.replace(config.settings, image_monthly_budget_usd=0.01))
    with pytest.raises(ImageGenerationError, match="ngân sách"):
        usage.reserve(8)


# ------------------------------------------------------------------ PPTX integration

def test_pptx_inserts_illustrations_by_slide_id(tmp_path):
    pres = outline_with_prompts()
    imgs = {}
    for sid in ("slide-01", "slide-02", "slide-04"):
        p = tmp_path / f"{sid}.jpg"
        p.write_bytes(jpeg())
        imgs[sid] = p
    imgs["slide-07"] = tmp_path / "missing.jpg"          # missing file is ignored, not fatal
    apply_illustrations(pres, {"slide-01", "slide-02", "slide-04"})
    out = tmp_path / "deck.pptx"
    build_pptx(pres, out, "professional", [], illustrations=imgs)
    report = validate_pptx(out, expected_slides=len(pres.slides))
    assert report.ok, report.errors
    assert report.pictures == 3 and not report.warnings
    prs = Pptx(str(out))
    texts = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "Ảnh minh họa tạo bởi AI" in texts and "Mục tiêu bài học" in texts   # text stays editable
    for idx in (0, 1, 3):
        pics = [sh for sh in prs.slides[idx].shapes if sh.shape_type == 13]
        assert len(pics) == 1
        pic = pics[0]
        ratio = (pic.width / pic.height) / (1 - pic.crop_left - pic.crop_right) * (1 - pic.crop_top - pic.crop_bottom)
        assert ratio == pytest.approx(1344 / 768, rel=0.02)   # not stretched
        # picture does not overlap any text box
        for sh in prs.slides[idx].shapes:
            if sh.has_text_frame and sh.text_frame.text.strip() and sh is not pic:
                overlap_x = min(sh.left + sh.width, pic.left + pic.width) - max(sh.left, pic.left)
                overlap_y = min(sh.top + sh.height, pic.top + pic.height) - max(sh.top, pic.top)
                assert overlap_x <= 0 or overlap_y <= 0, sh.text_frame.text[:30]


# ------------------------------------------------------------------ job pipeline + API

@pytest.fixture
def api_client(monkeypatch, image_settings):
    from fastapi.testclient import TestClient

    from app import main, pipeline
    from app.transcribe import Segment, Transcript

    def fake_transcribe(audio, language=None, progress=None):
        segs = [Segment(i * 1.0, i * 1.0 + 1, f"Câu thử nghiệm số {i} về quản lý thời gian.") for i in range(6)]
        return Transcript(segs, "vi", 6.0)

    monkeypatch.setattr(pipeline, "transcribe", fake_transcribe)
    monkeypatch.setattr(main, "settings", dataclasses.replace(main.settings, ai_images_enabled=True,
                                                              image_provider="mock", admin_token="adm1n"))
    with TestClient(main.app) as c:
        yield c


def _wait(client, job_id, timeout=60):
    end = time.time() + timeout
    while time.time() < end:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] in ("done", "failed"):
            return job
        time.sleep(0.2)
    raise AssertionError("job did not finish")


def _upload(client, mp4, **extra):
    with mp4.open("rb") as fh:
        return client.post("/api/jobs", files={"file": ("v.mp4", fh, "video/mp4")},
                           data={"language": "vi", **extra})


def test_config_exposes_ai_image_settings(api_client):
    cfg = api_client.get("/api/config").json()["ai_images"]
    assert cfg["enabled"] and cfg["available"] and cfg["placeholder"] and cfg["max_per_presentation"] == 8
    assert "MODAL" not in json.dumps(api_client.get("/api/config").json()).upper()  # no secrets to the browser


def test_job_with_ai_images_end_to_end(api_client, sample_mp4):
    r = _upload(api_client, sample_mp4, ai_images="true", image_density="standard")
    assert r.status_code == 202, r.text
    job_id = r.json()["job_id"]
    job = _wait(api_client, job_id)
    assert job["status"] == "done", job["error"]
    assert {"planning_illustrations", "generating_images", "generating"} <= set(job["stage_times"])
    il = job["result"]["illustrations"]
    assert il["requested"] == 1 and il["generated"] == 1 and il["real_model"] is False
    assert any("giả lập" in n for n in job["result"]["notices"])
    items = api_client.get(f"/api/jobs/{job_id}/illustrations").json()["items"]
    assert items[0]["slide_id"] == "slide-01" and items[0]["has_image"]
    img = api_client.get(f"/api/jobs/{job_id}/illustrations/slide-01.jpg")
    assert img.status_code == 200 and img.headers["content-type"] == "image/jpeg"
    assert api_client.get(f"/api/jobs/{job_id}/illustrations/..%2Fjob.jpg").status_code == 404
    assert api_client.get(f"/api/jobs/{job_id}/download").status_code == 200

    # Regenerate a single image: new seed, deck rebuilt, transcript/outline reused.
    before = img.content
    r = api_client.post(f"/api/jobs/{job_id}/illustrations/slide-01/regenerate")
    assert r.status_code == 202, r.text
    job = _wait(api_client, job_id)
    assert job["status"] == "done" and job["options"]["illustration_variants"] == {"slide-01": 1}
    assert api_client.get(f"/api/jobs/{job_id}/illustrations/slide-01.jpg").content != before
    assert api_client.post(f"/api/jobs/{job_id}/illustrations/slide-99/regenerate").status_code == 404


def test_job_without_ai_images_is_unchanged(api_client, sample_mp4):
    job = _wait(api_client, _upload(api_client, sample_mp4).json()["job_id"])
    assert job["status"] == "done"
    assert "generating_images" not in job["stage_times"] and job["result"]["illustrations"] is None


def test_requesting_images_when_unavailable_is_refused(api_client, sample_mp4):
    usage.set_kill_switch(True)
    try:
        r = _upload(api_client, sample_mp4, ai_images="true")
        assert r.status_code == 409 and "tạm dừng" in r.json()["detail"]
    finally:
        usage.set_kill_switch(False)


def test_failed_images_fail_the_job_then_retry_reuses_images(api_client, sample_mp4, monkeypatch):
    from app import pipeline

    flaky = MockImageProvider(fail_codes={"slide-01": ["invalid_params"]})
    monkeypatch.setattr(pipeline, "get_image_provider", lambda: flaky)
    job_id = _upload(api_client, sample_mp4, ai_images="true").json()["job_id"]
    job = _wait(api_client, job_id)
    assert job["status"] == "failed" and "ảnh minh họa" in job["error"]
    assert api_client.post(f"/api/jobs/{job_id}/retry").status_code == 202
    job = _wait(api_client, job_id)
    assert job["status"] == "done" and job["result"]["illustrations"]["generated"] == 1


def test_skip_policy_keeps_deck_with_explicit_notice(api_client, sample_mp4, monkeypatch):
    from app import pipeline

    monkeypatch.setattr(config, "settings", dataclasses.replace(config.settings, image_on_failure="skip"))
    monkeypatch.setattr(pipeline, "get_image_provider",
                        lambda: MockImageProvider(fail_codes={"slide-01": ["invalid_params"]}))
    job = _wait(api_client, _upload(api_client, sample_mp4, ai_images="true").json()["job_id"])
    assert job["status"] == "done"
    assert job["result"]["illustrations"]["generated"] == 0
    assert any("không tạo được" in n for n in job["result"]["notices"])


def test_admin_endpoints_require_token(api_client):
    assert api_client.get("/api/admin/images/usage").status_code == 401
    assert api_client.get("/api/admin/images/usage", headers={"X-Admin-Token": "nope"}).status_code == 401
    h = {"X-Admin-Token": "adm1n"}
    assert api_client.get("/api/admin/images/usage", headers=h).json()["limits"]["per_presentation"] == 8
    assert api_client.post("/api/admin/images/kill-switch", json={"active": True}, headers=h).json()["kill_switch"]
    assert not api_client.get("/api/config").json()["ai_images"]["available"]
    api_client.post("/api/admin/images/kill-switch", json={"active": False}, headers=h)


def test_cancel_during_gpu_batch_stops_and_cancels_pending_work(tmp_path, monkeypatch):
    provider = MockImageProvider(pending_polls=10_000)
    cancelled = []
    monkeypatch.setattr(provider, "cancel", cancelled.append)
    polls = {"n": 0}

    def should_cancel():
        polls["n"] += 1
        return polls["n"] > 2  # the user presses "cancel" while images are still pending

    service = IllustrationService(tmp_path / "j", provider, sleep=lambda _: None, should_cancel=should_cancel)
    with pytest.raises(ImageGenerationError) as e:
        service.run(_plan(2))
    assert e.value.code == "cancelled" and len(cancelled) == 2
    # Cancelled entries are re-planned, so a retry submits them again instead of polling dead calls.
    assert {e["state"] for e in service.load_manifest().values()} == {"planned"}
